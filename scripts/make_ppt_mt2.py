from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

BG=RGBColor(0x0D,0x1B,0x2A); AC=RGBColor(0x2E,0xCC,0x71)
A2=RGBColor(0x3A,0x9B,0xD5); WH=RGBColor(0xFF,0xFF,0xFF)
GR=RGBColor(0xB0,0xB8,0xC8); DC=RGBColor(0x16,0x28,0x3A)
YL=RGBColor(0xF3,0xC6,0x23); RD=RGBColor(0xE7,0x4C,0x3C)
IB=RGBColor(0x1E,0x35,0x4A)
W=Inches(13.33); H=Inches(7.5)

def mkprs():
    p=Presentation(); p.slide_width=W; p.slide_height=H; return p
def mks(p): return p.slides.add_slide(p.slide_layouts[6])
def bg(s):
    f=s.background.fill; f.solid(); f.fore_color.rgb=BG
def rc(s,l,t,w,h,c):
    sh=s.shapes.add_shape(1,l,t,w,h)
    sh.fill.solid(); sh.fill.fore_color.rgb=c; sh.line.fill.background()
def tx(s,text,l,t,w,h,sz=14,b=False,c=WH,al=PP_ALIGN.LEFT):
    bx=s.shapes.add_textbox(l,t,w,h); tf=bx.text_frame; tf.word_wrap=True
    p=tf.paragraphs[0]; p.alignment=al; r=p.add_run()
    r.text=text; r.font.size=Pt(sz); r.font.bold=b; r.font.color.rgb=c
def im(s,l,t,w,h,lb):
    sh=s.shapes.add_shape(1,l,t,w,h)
    sh.fill.solid(); sh.fill.fore_color.rgb=IB
    sh.line.color.rgb=A2; sh.line.width=Pt(1)
    tx(s,'[ '+lb+' ]',l,t+(h-Inches(0.4))/2,w,Inches(0.4),sz=10,c=A2,al=PP_ALIGN.CENTER)
def cd(s,l,t,w,h): rc(s,l,t,w,h,DC)
def bdg(s,lb):
    rc(s,Inches(0.4),Inches(0.27),Inches(3.8),Inches(0.44),AC)
    tx(s,lb,Inches(0.45),Inches(0.25),Inches(3.7),Inches(0.48),sz=13,b=True,c=BG,al=PP_ALIGN.CENTER)
def hd(s,t1,t2=None,t3=None):
    tx(s,t1,Inches(0.5),Inches(0.85),Inches(12.3),Inches(0.85),sz=28,b=True)
    if t2: tx(s,t2,Inches(0.5),Inches(1.65),Inches(5.5),Inches(0.42),sz=12,c=GR)
    if t3: tx(s,t3,Inches(6.2),Inches(1.65),Inches(6.8),Inches(0.42),sz=10.5,c=YL)
    rc(s,Inches(0.5),Inches(2.18),Inches(12.3),Inches(0.04),AC)

pr=mkprs()
s=mks(pr); bg(s)
bdg(s,"05  실험 결과")
hd(s,"Multi-task Ablation 결과",
   "EfficientNet-B3, full 전처리, 3-seed 평균",
   "* MT ablation은 val 최고 모델에서 수행 / 최종 best model(Swin)에 대한 FO/MT ablation은 추후 진행")

# 좌측 이미지 플레이스홀더
im(s, Inches(0.4), Inches(2.35), Inches(6.2), Inches(2.25),
   "fig_multitask_effect.png  (SO/SX/FO/MT val_acc)")
im(s, Inches(0.4), Inches(4.7),  Inches(6.2), Inches(2.4),
   "fig_female_vs_male_accuracy.png  (SO vs MT vs SX — test male/female)")

# 결과표
cd(s, Inches(6.85), Inches(2.35), Inches(6.05), Inches(2.8))
tx(s,"MT Ablation 결과표",
   Inches(7.0), Inches(2.45), Inches(5.7), Inches(0.42), sz=13, b=True, c=A2)

for h_txt,cx,cw in [("Mode",7.0,1.3),("Val Acc",8.45,1.1),("Sex Acc",9.65,1.0),("비고",10.75,1.95)]:
    tx(s,h_txt,Inches(cx),Inches(2.92),Inches(cw),Inches(0.38),sz=11,b=True,c=A2)

rows=[
    ("SO",   "84.0%", "-",     "기준선",  False),
    ("SX ★", "85.0%", "96.6%", "+1.0%p", True),
    ("FO",   "82.2%", "-",     "-1.8%p",  False),
    ("MT",   "83.2%", "94.3%", "-0.8%p",  False),
]
for i,(mode,val,sex,note,is_best) in enumerate(rows):
    ry = 3.38 + i*0.62
    if is_best:
        rc(s,Inches(6.95),Inches(ry-0.05),Inches(5.75),Inches(0.52),RGBColor(0x0A,0x22,0x12))
    clr=AC if is_best else WH
    for txt,cx,cw in [(mode,7.0,1.3),(val,8.45,1.1),(sex,9.65,1.0),(note,10.75,1.95)]:
        tx(s,txt,Inches(cx),Inches(ry),Inches(cw),Inches(0.48),sz=12,b=is_best,c=clr)

# 해석 박스 (수정된 버전)
cd(s, Inches(6.85), Inches(5.28), Inches(6.05), Inches(1.85))

# 왼쪽 해석
tx(s,"SX: val +1.0%p (N1 잠정) — test에서 역전 (SX 70.4% < SO 72.9%)",
   Inches(7.0), Inches(5.35), Inches(3.9), Inches(0.35), sz=10.5, c=WH)
tx(s,"-> val 과적합 가능성, 데이터 확장 후 재검증 필요",
   Inches(7.0), Inches(5.72), Inches(3.9), Inches(0.32), sz=10.5, c=GR)

tx(s,"하단 차트 (N1+N2 종합):",
   Inches(7.0), Inches(6.1), Inches(3.9), Inches(0.32), sz=10.5, b=True, c=YL)
tx(s,"SO: 수컷(74.5%) < 암컷(94.4%) — 형태 다형성으로 수컷이 더 어려움 (N2)",
   Inches(7.0), Inches(6.44), Inches(3.9), Inches(0.32), sz=10, c=WH)
tx(s,"SX/MT: 수컷 향상, 암컷 하락 — 암컷은 종간 형태차이 작아 보조신호 불리",
   Inches(7.0), Inches(6.78), Inches(3.9), Inches(0.32), sz=10, c=GR)

# 오른쪽 DINOv2
tx(s,"DINOv2 SX: 5.7% -> 완전 붕괴",
   Inches(11.1), Inches(5.35), Inches(1.65), Inches(0.35), sz=10, c=RD)
tx(s,"Self-supervised backbone",
   Inches(11.1), Inches(5.72), Inches(1.65), Inches(0.32), sz=10, c=GR)
tx(s,"multi-task gradient에 민감",
   Inches(11.1), Inches(6.06), Inches(1.65), Inches(0.32), sz=10, c=GR)

pr.save("beetledex_midterm_2026_MT수정2.pptx")
print("저장 완료")
